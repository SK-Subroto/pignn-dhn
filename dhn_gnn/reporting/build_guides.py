"""Build one step-by-step .docx guide per approach.

Each guide answers three questions for a single approach: how it works, which
source file does what and when it runs, and how to run it.

    python -m dhn_gnn.reporting.build_guides ["output directory"]

Numbers are read from results/report_data.json where available, so the guides
stay consistent with the measured run.
"""

import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from dhn_gnn import config

ACCENT = {"newton": RGBColor(0x0F, 0x76, 0x6E),
          "unrolled": RGBColor(0xB4, 0x53, 0x09),
          "initializer": RGBColor(0x1D, 0x4E, 0xD8)}
MUTED = RGBColor(0x55, 0x60, 0x63)
CODE = RGBColor(0x1F, 0x2A, 0x2E)


# ------------------------------------------------------------------ scaffolding
def _field(paragraph, instr):
    run = paragraph.add_run()
    for kind, payload in (("begin", None), ("instr", instr), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = payload
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def new_doc(accent):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.4)
    s.top_margin = s.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for name, size in (("Heading 1", 15.5), ("Heading 2", 12.5), ("Heading 3", 11)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = accent
        st.font.bold = True
        st.paragraph_format.space_before = Pt(14)
        st.paragraph_format.space_after = Pt(5)

    foot = s.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(foot, "PAGE")
    foot.runs[0].font.size = Pt(9)
    foot.runs[0].font.color.rgb = MUTED
    return doc


def para(doc, text, size=10.5, italic=False, bold=False, color=None, align=None, after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size); r.italic = italic; r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def rich(doc, parts, size=10.5, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    for item in parts:
        text, b, i = item if len(item) == 3 else (item[0], False, False)
        r = p.add_run(text)
        r.font.size = Pt(size); r.bold = b; r.italic = i
    return p


def code(doc, lines, indent=0.6):
    for line in lines:
        p = doc.add_paragraph()
        r = p.add_run(line)
        r.font.name = "Consolas"; r.font.size = Pt(9); r.font.color.rgb = CODE
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(indent)


def step(doc, number, title, body_parts, files=None):
    """A numbered step: bold heading, prose, and the code that carries it out."""
    p = doc.add_paragraph()
    r = p.add_run("Step " + str(number) + "   ")
    r.bold = True; r.font.size = Pt(10.5)
    r = p.add_run(title)
    r.bold = True; r.font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(10)
    rich(doc, body_parts)
    if files:
        q = doc.add_paragraph()
        r = q.add_run("runs:  ")
        r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = MUTED
        r = q.add_run(files)
        r.font.name = "Consolas"; r.font.size = Pt(8.5); r.font.color.rgb = MUTED
        q.paragraph_format.left_indent = Cm(0.6)
        q.paragraph_format.space_after = Pt(9)


def table(doc, headers, rows, number, cap, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True; r.font.size = Pt(8.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(8.5)
            if i == 0:
                r.font.name = "Consolas"
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w
    p = doc.add_paragraph()
    r = p.add_run("Table " + str(number) + ". " + cap)
    r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = MUTED
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    return t


def title_block(doc, accent, name, subtitle, tagline):
    para(doc, name, size=19, bold=True, color=accent,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=3)
    para(doc, subtitle, size=11.5, italic=True, color=MUTED,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    para(doc, tagline, size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)


SHARED_FILES = [
    ("config.py", "Paths, fluid constants, model defaults, and the train/test split.",
     "First, on every run."),
    ("physics/operators.py",
     "Builds the network once with PyDHN and extracts the fixed matrices: cycle matrix "
     "B, incidence A, pipe geometry.", "Once at start-up, then cached."),
    ("physics/darcy.py",
     "The pressure-drop model: Reynolds number, friction factor, phi and its "
     "derivative. Pure functions, no state.", "Inside every solver step."),
    ("solvers/physics_base.py",
     "DHNPhysicsBase. Holds the network matrices and wraps physics.py into edge_flow, "
     "pipe_dp, dp_der, residual and newton_step.", "Inherited by all three solvers."),
    ("datasets.py",
     "make_samples builds (mdot0, a_star, mdot_true) per hour; timestep_split chooses "
     "which hours are train and which are test.", "Before training and evaluation."),
    ("physics/pressure.py",
     "Reconstructs node pressure by least-squares integration of dp. Not part of the "
     "solve.", "During reporting only."),
    ("reporting/full_report.py",
     "Runs all three approaches on the same hours; writes prediction CSVs and figures.",
     "When building the comparison."),
]


def files_section(doc, specific, number):
    doc.add_heading(str(number) + "  Which file does what, and when", level=1)
    para(doc, "Paths are relative to dhn_gnn/. Files specific to this approach:", after=4)
    table(doc, ["File", "What it does", "When it runs"], specific, 1,
          "Files specific to this approach.", widths=[Cm(4.6), Cm(7.6), Cm(3.6)])
    para(doc, "Files shared with the other two approaches:", after=4)
    table(doc, ["File", "What it does", "When it runs"], SHARED_FILES, 2,
          "Shared infrastructure, identical for all three approaches. Sharing it is what "
          "makes the comparison fair: only the mechanism under study differs.",
          widths=[Cm(4.6), Cm(7.6), Cm(3.6)])


def results_section(doc, d, key, number, extra=None):
    doc.add_heading(str(number) + "  What to expect", level=1)
    if d:
        s = d["scores"][key]
        table(doc, ["Metric", "Value"], [
            ["solver steps, mean", "%.2f" % s["steps_mean"]],
            ["solver steps, median / max",
             "%.0f / %.0f" % (s["steps_median"], s["steps_max"])],
            ["wall clock", "%.0f ms per hour (CPU)" % s["ms_per_ts"]],
            ["flow MAE", "%.3e kg/s" % s["flow_mae"]],
            ["flow R2", "%.6f" % s["flow_r2"]],
            ["hours within tolerance", "%.1f %%" % s["pct_under_tol"]],
        ], 3, "Measured on %d held-out hours of 2022 at %.0f Pa tolerance."
              % (d["n_timesteps"], d["tol"]), widths=[Cm(7.5), Cm(8.3)])
    for block in (extra or []):
        rich(doc, block)


def footer_note(doc):
    p = doc.add_paragraph()
    r = p.add_run("Generated " + format(date.today(), "%d %B %Y") +
                  " from branch feat/learned-initializer. Measured values are read from "
                  "results/report_data.json, so re-running the report regenerates this "
                  "guide with current numbers.")
    r.font.size = Pt(8.5); r.italic = True; r.font.color.rgb = MUTED
    p.paragraph_format.space_before = Pt(18)


# ---------------------------------------------------------------------- newton
def guide_newton(d, out):
    doc = new_doc(ACCENT["newton"])
    title_block(doc, ACCENT["newton"], "The Newton Solver",
                "Pure physics, no machine learning",
                "How it works, which file does what, and how to run it")

    doc.add_heading("1  The idea in one paragraph", level=1)
    rich(doc, [("This solver contains ", False, False),
               ("no neural network and no learned parameters at all", True, False),
               (". It is a classical Newton-Raphson solve of the loop equations, written "
                "in PyTorch. It exists for two reasons: it is the control condition that "
                "every learned approach must beat, and it is the second half of the "
                "learned initializer, which only replaces its starting point.", False, False)])
    rich(doc, [("It matters because it is already very good. Deleting the neural network "
                "from the original architecture made the solver roughly ", False, False),
               ("nine times faster and more accurate", True, False),
               (". Any claim that machine learning helps has to be measured against this, "
                "not against the original model.", False, False)])

    doc.add_heading("2  How one hour is solved, step by step", level=1)
    step(doc, 1, "Build the network matrices",
         [("PyDHN constructs the network once. From it come the cycle matrix B (which "
           "loops exist), the incidence matrix A (which pipe touches which junction), and "
           "the pipe geometry. These never change, so they are built once and reused for "
           "every hour.", False, False)],
         "physics/operators.py -> build_operators()")
    step(doc, 2, "Prepare the hour",
         [("Read the mass flows for this hour. Split them into a part fixed by the "
           "boundary conditions (mdot0) and the part that circulates around the loops. "
           "Only the second part is unknown.", False, False)],
         "datasets.py -> make_samples()")
    step(doc, 3, "Start from zero",
         [("The starting guess is c = 0: assume nothing circulates. This is the ", False, False),
          ("cold start", False, True),
          (". It is honest but uninformed, and the residual begins around 3.6e4 Pa.",
           False, False)],
         "solvers/newton.py -> NewtonSolver.initial_guess()")
    step(doc, 4, "Measure how wrong it is",
         [("Compute the flow implied by the current guess, then the pressure drop in every "
           "pipe, then sum those drops around each loop. If the loops balanced perfectly "
           "the sum would be zero; whatever is left over is the residual, in pascals.",
           False, False)],
         "solvers/physics_base.py -> edge_flow(), pipe_dp(), residual()")
    step(doc, 5, "Work out the correction",
         [("Build the Jacobian J = B diag(dphi/dm) B-transpose. Because the problem was "
           "reduced to 12 loop unknowns, this matrix is only ", False, False),
          ("12 by 12", True, False),
          (" and solving it exactly costs almost nothing. Solving it exactly is what gives "
           "quadratic convergence: the error squares each time, so 100 Pa becomes 1 Pa "
           "becomes 0.0001 Pa.", False, False)],
         "solvers/physics_base.py -> newton_step(mode='full')")
    step(doc, 6, "Apply it, and check",
         [("Subtract the correction from the 12 loop flows and recompute the residual. If "
           "it is below the 50 Pa tolerance, stop. Otherwise return to step 4. On this "
           "network this loop runs about four times.", False, False)],
         "solvers/newton.py -> NewtonSolver.forward()")
    step(doc, 7, "Expand back to every pipe",
         [("The final 12 numbers are turned back into all 1,514 edge flows by "
           "mdot = mdot0 + Z c. Mass conservation is guaranteed by construction here, "
           "because A Z = 0 - it is never learned and cannot be violated.", False, False)],
         "solvers/physics_base.py -> edge_flow()")

    files_section(doc, [
        ("solvers/newton.py", "NewtonSolver. The solve loop, the cold start, and the "
         "stopping rule.", "Every hour."),
        ("solvers/physics_base.py", "newton_step(mode='full') builds and solves the 12x12 system.",
         "Every solver step."),
    ], 3)

    doc.add_heading("4  How to run it", level=1)
    rich(doc, [("There is ", False, False), ("no training step", True, False),
               (". This solver has no parameters, so there is nothing to fit and no "
                "checkpoint to load. It is constructed and used directly.", False, False)])
    para(doc, "Run it as part of the comparison, which reports it alongside the other two:",
         after=4)
    code(doc, ["python -m dhn_gnn.reporting.full_report --n-ts 400"])
    para(doc, "Or use it from Python on its own:", after=4)
    code(doc, [
        "from dhn_gnn.physics import operators as netops",
        "from dhn_gnn.solvers.newton import NewtonSolver",
        "from dhn_gnn.datasets import make_samples",
        "",
        "ops = netops.build_operators()",
        "model = NewtonSolver(ops, n_newton=8).float().eval()",
        "m0, a_star, mdot_true = make_samples(ops, [100])[0]",
        "mdot, residuals, c_hist = model(m0, tol=50.0)",
    ])
    para(doc, "Outputs written by the comparison run:", after=4)
    code(doc, ["results/predictions/newton/mass_flow.csv",
               "results/predictions/newton/delta_p.csv",
               "results/predictions/newton/pressure.csv"])

    results_section(doc, d, "newton", 5, extra=[
        [("The residual keeps falling after the tolerance is met - exact Newton overshoots "
          "the target rather than stopping precisely on it. That is why a faster solver "
          "can report a ", False, False), ("higher", False, True),
         (" final residual: it stopped sooner. Compare steps, not final residual.",
          False, False)],
    ])
    footer_note(doc)
    doc.save(out)
    return out


# -------------------------------------------------------------------- unrolled
def guide_unrolled(d, out):
    doc = new_doc(ACCENT["unrolled"])
    title_block(doc, ACCENT["unrolled"], "The Unrolled GNN Solver",
                "The original architecture: a graph network at every step",
                "How it works, which file does what, and how to run it")

    doc.add_heading("1  The idea in one paragraph", level=1)
    rich(doc, [("This was the project's first design. It runs a fixed number of solver "
                "steps - twenty - and at ", False, False), ("every one", False, True),
               (" of them a graph neural network looks at the whole network and proposes a "
                "correction to the 12 loop flows. That correction is added on top of a "
                "damped Newton step. Each of the twenty steps has its ", False, False),
               ("own separate head", True, False),
               (", so the model can in principle learn a different behaviour early in the "
                "solve than late in it.", False, False)])
    rich(doc, [("The heads are zero-initialised, which means an untrained model is exactly "
                "the physics solver. Training can only move away from that starting point - "
                "and on this dataset it never found anything better. This document "
                "describes the approach and records that result.", False, False)])

    doc.add_heading("2  How one hour is solved, step by step", level=1)
    para(doc, "Steps 1 and 2 are identical to the other approaches: build the network "
              "matrices, then prepare the hour. What follows is repeated twenty times.",
         after=6)
    step(doc, 3, "Start from zero",
         [("The loop flows start at c = 0, the same cold start the pure physics solver "
           "uses.", False, False)],
         "approaches/unrolled/model.py -> forward()")
    step(doc, 4, "Evaluate the physics at the current guess",
         [("Compute pipe pressure drops, their derivatives, and the loop residual - the "
           "same three quantities the pure Newton solver uses.", False, False)],
         "solvers/physics_base.py -> pipe_dp(), dp_der(), residual()")
    step(doc, 5, "Describe the network to the graph network",
         [("Build four numbers per junction: is it a boundary node, how much water is "
           "injected, the summed pressure drop, and the summed flow magnitude. These are "
           "projected into a 64-dimensional embedding.", False, False)],
         "approaches/unrolled/model.py -> node_proj")
    step(doc, 6, "Message passing with attention",
         [("Two rounds of attention. Each junction attends to its neighbours along the "
           "pipes, and each pipe adds a bias derived from its current hydraulic resistance, "
           "so the physics steers where the network pays attention. This is the expensive "
           "part, and it happens on ", False, False), ("every one of the twenty steps",
                                                        True, False), (".", False, False)],
         "solvers/attention.py -> EdgeBiasedAttention")
    step(doc, 7, "Collapse to the 12 loops and propose a correction",
         [("Junction embeddings are read out onto pipes, then summed onto loops. Step k's "
           "own head turns that into one number per loop, bounded by a tanh so it cannot "
           "explode.", False, False)],
         "approaches/unrolled/model.py -> edge_readout, heads[k]")
    step(doc, 8, "Combine with a damped Newton step",
         [("The update is dc = -0.5 x newton_step + 0.01 x tanh(head). The Newton part uses "
           "only the ", False, False), ("diagonal", False, True),
          (" of the Jacobian - it pretends the loops do not affect each other - which is "
           "why it needs the 0.5 damping and converges slowly.", False, False)],
         "solvers/physics_base.py -> newton_step(mode='diagonal')")
    step(doc, 9, "Repeat, then expand",
         [("Steps 4 to 8 run twenty times, then the final 12 loop flows expand back to all "
           "1,514 pipe flows. On this network the tolerance is reached after about "
           "sixteen steps.", False, False)],
         "approaches/unrolled/model.py -> forward()")

    doc.add_heading("3  How training works", level=1)
    rich(doc, [("Training uses a ", False, False), ("deep physics loss", True, False),
               (": every one of the twenty steps is scored on the residual it produces, "
                "with later steps weighted more heavily. Supervising only the final answer "
                "was tried first and did not work - the gradient had to travel through "
                "twenty steps of a quadratic operator and exploded.", False, False)])
    rich(doc, [("Even with deep supervision, the state is ", False, False),
               ("detached at each step", True, False),
               (" (truncated backpropagation), so gradients only flow within one step. "
                "Without this the observed gradient magnitude reached about 1e14.",
                False, False)])
    rich(doc, [("A practical warning. The learning rate matters enormously here. At the "
                "original lr=1e-3 with step_scale=0.5 the residual goes from 35 Pa to "
                "about 180,000 Pa ", False, False), ("within a single epoch", True, False),
               (". A head step of plus or minus 0.5 kg/s is enormous next to loop flows of "
                "3 to 5 kg/s. Use lr=1e-5 and step_scale=0.01.", False, False)])

    files_section(doc, [
        ("approaches/unrolled/model.py", "DHNUnrolledSolver. The twenty-step loop, the "
         "per-step heads, and the combined update.", "Every hour, twenty times."),
        ("solvers/attention.py", "EdgeBiasedAttention. Message passing between junctions, "
         "biased by pipe resistance.", "Twenty times per hour."),
        ("approaches/unrolled/train.py", "deep_physics_loss and the training loop with "
         "best-weight restore.", "During training only."),
        ("losses.py", "The earlier discounted physics loss, kept for the smoke test.",
         "Tests only."),
    ], 4)

    doc.add_heading("5  How to run it", level=1)
    para(doc, "Train, reproducing the original configuration (twenty steps, diagonal "
              "Newton):", after=4)
    code(doc, ["python -m dhn_gnn.cli train --arch unrolled --k 20 "
               "--newton-mode diagonal \\",
               "    --epochs 10 --n-train 60 --lr 1e-5 --step-scale 0.01 \\",
               "    --out results/model_unrolled_2022.pt"])
    para(doc, "Then evaluate:", after=4)
    code(doc, ["python -m dhn_gnn.cli evaluate --arch unrolled "
               "--ckpt results/model_unrolled_2022.pt"])
    para(doc, "Expect training to take roughly ten minutes on CPU. Each hour costs about "
              "1.2 seconds because attention runs twenty times per forward pass, and "
              "again on the backward pass.", after=8)

    results_section(doc, d, "unrolled", 6, extra=[
        [("An honest note on training. ", True, False),
         ("On this dataset, best-weight selection returned ", False, False),
         ("epoch zero, the initialisation", True, False),
         (". Every later epoch was worse on the monitored residual, at every stable "
          "learning rate tried. Because the heads are zero-initialised, that means the "
          "checkpoint is effectively the pure physics solver, carrying the cost of twenty "
          "attention passes without the benefit.", False, False)],
        [("This is a legitimate result to report rather than a failure to hide. The "
          "physics step was already at the accuracy the reference data supports, so the "
          "learned correction had no headroom to work in. Moving the same network from "
          "correcting to guessing is what made it useful - see the learned initializer "
          "guide.", False, False)],
    ])
    footer_note(doc)
    doc.save(out)
    return out


# ----------------------------------------------------------------- initializer
def guide_initializer(d, out):
    doc = new_doc(ACCENT["initializer"])
    title_block(doc, ACCENT["initializer"], "The Learned Initializer",
                "One graph-network pass for the guess, exact Newton for the answer",
                "How it works, which file does what, and how to run it")

    doc.add_heading("1  The idea in one paragraph", level=1)
    rich(doc, [("Neural networks are good at being ", False, False),
               ("approximately right, instantly", False, True),
               (". Newton's method is good at being ", False, False),
               ("exactly right, given a decent starting point", False, True),
               (". This approach pairs them so each does only what it is good at: the "
                "graph network runs ", False, False), ("once", True, False),
               (" and predicts the 12 loop flows, then parameter-free Newton corrects "
                "them. The network never competes with Newton on precision.", False, False)])
    rich(doc, [("The deeper reason this design was chosen is not step count. PyDHN also "
                "gets a good starting point - by reusing the previous hour - but that "
                "forces it to solve the year strictly in order. A guess computed from the "
                "boundary conditions alone carries ", False, False),
               ("no dependency on any other hour", True, False),
               (", so all 8,760 hours are independent problems that could be solved at "
                "the same time.", False, False)])

    doc.add_heading("2  How one hour is solved, step by step", level=1)
    para(doc, "Steps 1 and 2 are identical to the other approaches: build the network "
              "matrices once, then prepare the hour.", after=6)
    step(doc, 3, "Describe the starting state to the network",
         [("Evaluate the physics at the boundary-feasible reference flow: pipe pressure "
           "drops and their derivatives. Build four numbers per junction - boundary flag, "
           "injection, summed pressure drop, summed flow magnitude - and project them into "
           "a 64-dimensional embedding.", False, False)],
         "approaches/initializer/model.py -> initial_guess()")
    step(doc, 4, "Two rounds of attention",
         [("Each junction exchanges information with its neighbours along the pipes. Each "
           "pipe contributes a bias derived from its hydraulic resistance, so the physics "
           "guides the message passing. This runs ", False, False),
          ("exactly once per hour", True, False),
          (" - the single most important difference from the original architecture.",
           False, False)],
         "solvers/attention.py -> EdgeBiasedAttention")
    step(doc, 5, "Collapse to 12 numbers",
         [("Junction embeddings are read out onto pipes, then averaged onto loops. "
           "Averaging rather than summing matters: a loop spans hundreds of pipes, and a "
           "plain sum produced features about a hundred times larger than the rest of the "
           "network expects, which made early predictions wildly off-scale.", False, False)],
         "approaches/initializer/model.py -> loop_norm, head")
    step(doc, 6, "The guess",
         [("The head outputs one number per loop: the predicted circulation. Its output "
           "layer is zero-initialised, so an untrained model predicts zero and degrades "
           "gracefully to the cold start. After training the guess is about 0.06 kg/s from "
           "the truth, roughly a hundred times closer than starting from zero.",
           False, False)],
         "approaches/initializer/model.py -> head")
    step(doc, 7, "Exact Newton finishes the job",
         [("From here the code is the pure physics solver, inherited unchanged: measure "
           "the residual, build the 12x12 Jacobian, solve it exactly, apply the "
           "correction, stop when below 50 Pa. ", False, False),
          ("No learned parameters are involved beyond this point.", True, False)],
         "solvers/newton.py -> NewtonSolver.forward()")
    step(doc, 8, "Expand back to every pipe",
         [("The final 12 loop flows become all 1,514 edge flows. Mass conservation holds "
           "by construction.", False, False)],
         "solvers/physics_base.py -> edge_flow()")

    doc.add_heading("3  How training works, and why it is easy", level=1)
    rich(doc, [("The correct answer for the 12 loop flows can be computed directly from the "
                "reference data: a_star is the loop-space coordinate of the true flow. So "
                "training is ", False, False),
               ("plain regression on 12 numbers", True, False),
               (" - predict a_star, minimise mean squared error. Nothing else.",
                False, False)])
    rich(doc, [("This is why it trains where the unrolled architecture did not. There is ",
                False, False), ("no solver in the training loop at all", True, False),
               (": no gradients through Newton, no unrolling, no truncated "
                "backpropagation, no per-step heads competing for gradient. Training runs "
                "zero Newton steps.", False, False)])
    rich(doc, [("Model selection uses the guess error, not the polished residual. Newton "
                "hides a mediocre guess by working harder, so scoring after the polish "
                "would measure the physics rather than the thing being trained.",
                False, False)])

    files_section(doc, [
        ("approaches/initializer/model.py", "DHNInitializerSolver. The attention stack and the head "
         "that predicts the 12 loop flows. Everything else is inherited.",
         "Once per hour."),
        ("solvers/newton.py", "The parent class. Supplies the Newton polish, unchanged.",
         "After the guess."),
        ("solvers/attention.py", "EdgeBiasedAttention. Message passing between junctions.",
         "Once per hour."),
        ("approaches/initializer/train.py", "train_initializer. Regression on a_star with "
         "best-weight restore.", "During training only."),
        ("checkpoints.py", "Saves weights together with the hyperparameters "
         "needed to rebuild the model.", "End of training, start of evaluation."),
    ], 4)

    doc.add_heading("5  How to run it", level=1)
    para(doc, "Train. About seven minutes on CPU for the settings below:", after=4)
    code(doc, ["python -m dhn_gnn.cli train --arch initializer \\",
               "    --epochs 40 --n-train 200 --lr 1e-3"])
    para(doc, "Evaluate on the held-out hours. This writes the prediction CSVs, the "
              "metrics summary and the figure:", after=4)
    code(doc, ["python -m dhn_gnn.cli evaluate --arch initializer"])
    para(doc, "Compare against the other two approaches and PyDHN:", after=4)
    code(doc, ["python -m dhn_gnn.reporting.full_report --n-ts 400",
               "python -m dhn_gnn.reporting.build_docx"])
    para(doc, "Useful flags:", after=4)
    table(doc, ["Flag", "Meaning", "Default"], [
        ("--n-train", "How many hours to fit on. Sampled evenly across the year.", "200"),
        ("--epochs", "Training passes over those hours.", "40"),
        ("--lr", "Adam learning rate. Far less delicate here than for the unrolled model.",
         "1e-3"),
        ("--out", "Checkpoint path. Point smoke runs elsewhere so a one-epoch test cannot "
         "overwrite a real model.", "results/model_init.pt"),
        ("--device", "cpu or cuda. See the note below before using cuda.", "cpu"),
    ], 3, "Command-line flags for training.", widths=[Cm(2.6), Cm(9.8), Cm(3.4)])

    results_section(doc, d, "initializer", 6, extra=[
        [("Where the gain comes from. ", True, False),
         ("The learned component contributes exactly one thing: a better starting point. "
          "Cold start begins around 3.6e4 Pa; the learned guess begins around 3.3e2 Pa. "
          "That saves roughly two Newton steps.", False, False)],
        [("A caveat worth stating to anyone reading the results. ", True, False),
         ("Those saved steps are cheap, so in serial wall-clock the gain over pure Newton "
          "is small. The gain that matters is structural: no dependency on the previous "
          "hour, so the year parallelises. Batching is what converts that into wall-clock, "
          "and it is not yet implemented.", False, False)],
        [("On GPU. ", True, False),
         ("Measured on a GTX 1060, inference took 104.8 ms against 44.0 ms on CPU, with "
          "training a tie. At one hour per forward pass the tensors are too small to cover "
          "kernel-launch overhead, so cpu is the default. Batching is the prerequisite for "
          "a GPU to help.", False, False)],
    ])
    footer_note(doc)
    doc.save(out)
    return out


# ------------------------------------------------------------------------ main
def main():
    outdir = Path(sys.argv[1]) if len(sys.argv) > 1 else config.RESULTS_DIR / "report"
    outdir.mkdir(parents=True, exist_ok=True)

    data_path = config.RESULTS_DIR / "report_data.json"
    d = json.loads(data_path.read_text(encoding="utf-8")) if data_path.exists() else None
    if d is None:
        print("warning: results/report_data.json not found; guides will omit measurements")

    made = [
        guide_newton(d, outdir / "1 - Newton solver (pure physics).docx"),
        guide_unrolled(d, outdir / "2 - Unrolled GNN solver (original).docx"),
        guide_initializer(d, outdir / "3 - Learned initializer (current).docx"),
    ]
    for p in made:
        print("wrote %s  (%.0f KB)" % (p, p.stat().st_size / 1024))


if __name__ == "__main__":
    main()
